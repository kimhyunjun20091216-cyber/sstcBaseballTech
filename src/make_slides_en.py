"""Presentation (pptx) generator — English deck for KSEF/ISEF, ~15 slides.

Standard font (Calibri), embeds result figures, 16:9. Mirrors src/make_slides.py.

Usage: python src/make_slides_en.py -o paper/PitchingAnalysis_Slides.pptx
"""

from __future__ import annotations

import argparse
import os

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

FONT = "Calibri"
NAVY = RGBColor(0x1F, 0x4E, 0x79)
LIGHT = RGBColor(0xEB, 0xF1, 0xF8)
GRAY = RGBColor(0x44, 0x44, 0x44)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLUEISH = RGBColor(0xBD, 0xD3, 0xEC)

FIG = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "paper", "figs")


def _font(run, size, bold=False, color=GRAY):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_title_slide(prs, title, subtitle, meta):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, NAVY)
    box = s.shapes.add_textbox(Inches(0.7), Inches(2.0), Inches(11.9), Inches(3.2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = title
    _font(r, 28, bold=True, color=WHITE)
    p2 = tf.add_paragraph(); p2.space_before = Pt(14)
    r = p2.add_run(); r.text = subtitle; _font(r, 15, color=BLUEISH)
    p3 = tf.add_paragraph(); p3.space_before = Pt(24)
    r = p3.add_run(); r.text = meta; _font(r, 13, color=BLUEISH)
    return s


def _header(slide, title, idx=None):
    bar = slide.shapes.add_shape(1, 0, 0, Emu(int(13.333 * 914400)), Inches(0.95))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    tf = bar.text_frame; tf.margin_left = Inches(0.5); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; r = p.add_run(); r.text = title
    _font(r, 22, bold=True, color=WHITE)
    if idx is not None:
        nb = slide.shapes.add_textbox(Inches(12.2), Inches(0.25), Inches(0.9), Inches(0.5))
        rp = nb.text_frame.paragraphs[0]; rp.alignment = PP_ALIGN.RIGHT
        rr = rp.add_run(); rr.text = str(idx); _font(rr, 12, color=BLUEISH)


def add_bullets(prs, title, bullets, idx=None):
    s = prs.slides.add_slide(prs.slide_layouts[6]); _bg(s, WHITE)
    _header(s, title, idx)
    box = s.shapes.add_textbox(Inches(0.7), Inches(1.3), Inches(11.9), Inches(5.9))
    tf = box.text_frame; tf.word_wrap = True
    for k, b in enumerate(bullets):
        level, text = (b if isinstance(b, tuple) else (0, b))
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.level = level; p.space_after = Pt(9)
        bullet = "•  " if level == 0 else "–  "
        r = p.add_run(); r.text = bullet + text
        _font(r, 18 if level == 0 else 15, bold=(level == 0), color=(NAVY if level == 0 else GRAY))
    return s


def add_image_slide(prs, title, img, caption, idx=None, bullets=None):
    s = prs.slides.add_slide(prs.slide_layouts[6]); _bg(s, WHITE)
    _header(s, title, idx)
    path = os.path.join(FIG, img)
    if bullets:
        if os.path.exists(path):
            s.shapes.add_picture(path, Inches(0.5), Inches(1.3), height=Inches(5.4))
        box = s.shapes.add_textbox(Inches(7.7), Inches(1.5), Inches(5.2), Inches(5.4))
        tf = box.text_frame; tf.word_wrap = True
        for k, b in enumerate(bullets):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.space_after = Pt(10)
            r = p.add_run(); r.text = "•  " + b; _font(r, 15, color=GRAY)
    else:
        if os.path.exists(path):
            s.shapes.add_picture(path, Inches(2.6), Inches(1.2), height=Inches(5.2))
    if caption:
        cap = s.shapes.add_textbox(Inches(0.5), Inches(6.85), Inches(12.3), Inches(0.5))
        rp = cap.text_frame.paragraphs[0]; rp.alignment = PP_ALIGN.CENTER
        r = rp.add_run(); r.text = caption; _font(r, 11, color=GRAY); r.font.italic = True
    return s


def add_table_slide(prs, title, header, rows, idx=None, note=None):
    s = prs.slides.add_slide(prs.slide_layouts[6]); _bg(s, WHITE)
    _header(s, title, idx)
    nrows, ncols = len(rows) + 1, len(header)
    gt = s.shapes.add_table(nrows, ncols, Inches(0.5), Inches(1.35),
                            Inches(12.3), Inches(0.4 * nrows)).table
    for j, h in enumerate(header):
        c = gt.cell(0, j); c.text = ""
        r = c.text_frame.paragraphs[0].add_run(); r.text = h
        _font(r, 12, bold=True, color=WHITE)
        c.fill.solid(); c.fill.fore_color.rgb = NAVY
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            c = gt.cell(i, j); c.text = ""
            r = c.text_frame.paragraphs[0].add_run(); r.text = val
            _font(r, 11, bold=(j == 0), color=GRAY)
            c.fill.solid(); c.fill.fore_color.rgb = LIGHT if i % 2 else WHITE
    if note:
        nb = s.shapes.add_textbox(Inches(0.5), Inches(6.9), Inches(12.3), Inches(0.5))
        r = nb.text_frame.paragraphs[0].add_run(); r.text = note
        _font(r, 11, color=GRAY); r.font.italic = True
    return s


def build(out_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1. Title
    add_title_slide(
        prs,
        "Markerless Baseball Pitching Analysis\nfrom a Single Smartphone Video",
        "A student-reproducible methodology for validating low-cost 2D pose estimation",
        "Sports Technology Club · Hyunjun Kim  |  KSEF 2026 / ISEF 2027",
    )

    # 2. Background
    add_bullets(prs, "Background — The Accessibility Gap", [
        "Pitching = the fastest human joint motion (shoulder ~7,000°/s), elbow varus torque ~64 N·m → high injury risk",
        "Standard analysis = optical marker mocap / multi-camera (KinaTrax, Hawk-Eye, Theia3D)",
        (1, "Requires costly hardware, controlled labs, expert staff"),
        "Result: students, youth, and amateurs who most need analysis cannot access it",
        "→ Can a single low-cost smartphone close this gap?",
    ], idx=2)

    # 3. Objectives
    add_bullets(prs, "Objectives and Research Questions", [
        "Initial requirement: generate a joint skeleton ('stick figure') on each of two camera angles",
        "RQ1. Stable pose estimation from one smartphone under different conditions (daylight side / backlit front)?",
        "RQ2. Can pitch phases and key joint kinematics be computed automatically?",
        "RQ3. Without ground-truth mocap, how do we validate the reliability of low-cost 2D measurement?",
    ], idx=3)

    # 4. Related work
    add_table_slide(prs, "Related Work Comparison",
        ["System", "Cameras", "Cost", "2D/3D", "Target", "Open"],
        [
            ["This work", "1 (smartphone)", "Consumer", "2D", "Student/amateur", "Yes"],
            ["pitchAI", "1", "Commercial", "3D", "Lab/pro", "No"],
            ["Fleisig 2024", "9", "Laboratory", "3D", "Pro", "No"],
            ["Aguinaldo 2025", "Multiple", "Very high", "3D", "Stadium", "No"],
            ["PitcherNet", "1 (broadcast)", "High compute", "3D", "Pro analytics", "No"],
            ["OpenCap/Pose2Sim", "2+", "Low–moderate", "3D", "Researchers", "Yes"],
        ], idx=4,
        note="2025 J Sports Sci: \"no current research exists on single-camera markerless feasibility for pitching analysis\"")

    # 5. Novelty
    add_bullets(prs, "Novelty — Accessibility, Not Accuracy", [
        "We do NOT claim higher accuracy than lab systems (marker mocap remains the standard)",
        "Contribution 1 — Accessibility: pipeline runs on a single consumer smartphone",
        "Contribution 2 — Cross-engine validation: two independent engines' agreement as a reliability proxy",
        "Contribution 3 — Reproducible methodology: Bland-Altman · ICC · RMSE",
        "No prior study applies this combination to baseball pitching",
    ], idx=5)

    # 6. Architecture
    add_bullets(prs, "System Architecture", [
        "Input video → preprocessing (optional CLAHE backlight correction)",
        (1, "→ pose engine (MediaPipe / YOLO11) → COCO 17 joints"),
        (1, "→ One-Euro smoothing (jitter removal) → skeleton overlay"),
        (1, "→ automatic phase segmentation → joint kinematics → validation stats"),
        "Entirely Python 3.11 + open source — reproducible by students",
    ], idx=6)

    # 7. Methods - pose/smoothing
    add_bullets(prs, "Methods — Pose Estimation and Smoothing", [
        "Two engines unified under one convention (COCO 17 joints) → swappable & comparable",
        (1, "MediaPipe BlazePose (Heavy): 33 landmarks, real-time on CPU"),
        (1, "YOLO11-pose: 17 joints, backup for backlit/small subjects"),
        "One-Euro filter: speed-adaptive → minimal lag at release, tremor suppressed when static",
        "In multi-person frames, auto-select the single pitcher (max joint-confidence sum)",
    ], idx=7)

    # 8. Methods - kinematics
    add_bullets(prs, "Methods — Kinematics and Phase Segmentation", [
        "Joint angle = interior angle of three points; elbow/knee as flexion (180° − interior)",
        "Trunk tilt = angle of pelvis→shoulder vector relative to vertical",
        "Automatic phase detection:",
        (1, "peak leg lift · foot contact · release (max wrist speed)"),
        "Throwing arm & lead leg auto-determined from max wrist speed",
    ], idx=8)

    # 9. Methods - validation
    add_bullets(prs, "Methods — Measurement-Reliability Validation", [
        "CV angles vs ground truth (two trained annotators, manual)",
        "Bland-Altman: bias + 95% limits of agreement",
        "ICC(2,1): two-way random, absolute agreement + 95% CI (Shrout & Fleiss 1979)",
        "RMSE · MAE · Pearson r; test-retest / inter- & intra-rater reliability",
        "When ground truth is absent → cross-engine agreement as proxy",
    ], idx=9)

    # 10. Result 1
    add_image_slide(prs, "Result 1 — 100% Detection, Both Angles (RQ1)",
        "fig_side_overlay.png", "Figure. Skeleton overlay on the side-view (daylight) video", idx=10,
        bullets=[
            "Side (daylight): 122/122 = 100%",
            "Front (backlit, small): 197/197 = 100%",
            "Anatomically plausible tracking even in the hard backlit front view",
            "Requirement (two-angle stick figures) satisfied",
        ])

    # 11. Result 2
    add_image_slide(prs, "Result 2 — Automatic Kinematics & Phases (RQ2)",
        "fig_kinematics.png", "Figure. Pitch-kinematics time series with auto phase markers", idx=11,
        bullets=[
            "Phases: leg lift (1.17s) → contact (2.10s) → release (2.33s)",
            "Peak wrist speed (~2,800 px/s) aligns exactly with release",
            "Lead knee 118° flexion → extension (lead-leg block)",
            "Trunk forward tilt 38° at release (within normal 28–42°)",
        ])

    # 12. Result 3
    add_bullets(prs, "Result 3 — Cross-Engine & Statistical Validation (RQ3)", [
        "Two engines localize the same points within 6–7% of torso length",
        (1, "Side 13.5px/7%, front 7.2px/6% → cross-validation feasible"),
        "MediaPipe: marginally lower jitter + ~7× faster → chosen as default engine",
        "ICC(2,1) implementation matches Shrout & Fleiss example (0.290) exactly ✓",
        "End-to-end validation on synthetic GT works → ready for real annotations",
    ], idx=12)

    # 13. Discussion
    add_bullets(prs, "Discussion — Limitations and Significance", [
        "Limitations:",
        (1, "2D single camera → shoulder rotation unmeasurable (deliberately omitted)"),
        (1, "30 fps undersamples release → 240 fps needed for formal capture"),
        (1, "No ground truth / small sample → absolute accuracy pending validation"),
        "Significance: an accessible pipeline + a scientifically honest, reproducible validation methodology",
    ], idx=13)

    # 14. Conclusion
    add_bullets(prs, "Conclusion and Future Work", [
        "Single smartphone → 100% detection & automatic kinematics; value = accessibility + rigor",
        "Future 1 — IRB-approved 240 fps, two-camera multi-pitcher dataset",
        "Future 2 — establish per-variable valid ranges vs manual-annotation ground truth",
        "Future 3 — sync force-plate GRF with video release → quantify 'lower body → ball'",
        "Future 4 — extend to multi-camera 3D (Pose2Sim / OpenCap)",
    ], idx=14)

    # 15. Summary
    add_bullets(prs, "Summary", [
        "Requirement (two-angle stick figures) → extended into a full research study",
        "Accessibility: single smartphone, open source, student-reproducible",
        "Rigor: cross-engine validation + Bland-Altman/ICC/RMSE (ICC verified)",
        "Honesty: 2D / 30 fps / shoulder-rotation limitations stated explicitly",
        "Thank you.",
    ], idx=15)

    prs.save(out_path)
    n = len(prs.slides.__iter__.__self__._sldIdLst)
    print(f"[pptx] saved: {out_path}  ({n} slides)")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("-o", "--output", default="paper/PitchingAnalysis_Slides.pptx")
    args = ap.parse_args()
    build(args.output)


if __name__ == "__main__":
    main()
