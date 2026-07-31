"""연구 포스터(pptx) 생성 — KSEF/ISEF 스타일 대형 단일 슬라이드.

48"x36" 가로, 3열 레이아웃, 한글 폰트(맑은 고딕), 결과 이미지 임베드.

사용: python src/make_poster.py -o paper/투구분석_포스터.pptx
"""

from __future__ import annotations

import argparse
import os

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

KOR = "맑은 고딕"
NAVY = RGBColor(0x1F, 0x4E, 0x79)
BLUE = RGBColor(0x2E, 0x6D, 0xA4)
LIGHT = RGBColor(0xEB, 0xF1, 0xF8)
PANEL = RGBColor(0xF6, 0xF9, 0xFC)
GRAY = RGBColor(0x33, 0x33, 0x33)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0xC0, 0x39, 0x2B)

FIG = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "paper", "figs")

# 포스터 치수(인치).
W, H = 48, 36
MARGIN = 1.2
GUTTER = 0.8
HEADER_H = 4.5
COL_TOP = MARGIN + HEADER_H + 0.6
COL_W = (W - 2 * MARGIN - 2 * GUTTER) / 3
COL_X = [MARGIN, MARGIN + COL_W + GUTTER, MARGIN + 2 * (COL_W + GUTTER)]


def _kfont(run, size, bold=False, color=GRAY):
    run.font.name = KOR
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.makeelement(qn("a:ea"), {"typeface": KOR})
    rPr.append(ea)


def _rect(slide, x, y, w, h, color, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = color
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(1)
    shp.shadow.inherit = False
    return shp


def _text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
    """runs: list of paragraphs; each paragraph = list of (text, size, bold, color)."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = wrap; tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.15); tf.margin_right = Inches(0.15)
    tf.margin_top = Inches(0.08); tf.margin_bottom = Inches(0.08)
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(6)
        for (txt, size, bold, color) in para:
            r = p.add_run(); r.text = txt; _kfont(r, size, bold, color)
    return box


def _panel(slide, col, y, h, header):
    """섹션 패널: 헤더 바 + 본문 배경. 반환: 본문 시작 y."""
    x = COL_X[col]
    _rect(slide, x, y, COL_W, 0.95, NAVY)
    _text(slide, x + 0.1, y + 0.05, COL_W - 0.2, 0.85,
          [[(header, 30, True, WHITE)]], anchor=MSO_ANCHOR.MIDDLE)
    _rect(slide, x, y + 0.95, COL_W, h - 0.95, PANEL, line=RGBColor(0xD5, 0xDE, 0xE8))
    return y + 1.15


def _bullets(items, size=20, color=GRAY):
    out = []
    for it in items:
        if isinstance(it, tuple):
            lvl, txt = it
            prefix = "   – " if lvl else "• "
            out.append([(prefix + txt, size - (2 if lvl else 0), False, color)])
        else:
            out.append([("• " + it, size, False, color)])
    return out


def _mini_table(slide, x, y, w, header, rows, cell_h=0.62, fs=15):
    n = len(rows) + 1
    tbl = slide.shapes.add_table(n, len(header), Inches(x), Inches(y),
                                 Inches(w), Inches(cell_h * n)).table
    for j, htxt in enumerate(header):
        c = tbl.cell(0, j); c.text = ""
        c.fill.solid(); c.fill.fore_color.rgb = NAVY
        rr = c.text_frame.paragraphs[0].add_run(); rr.text = htxt
        _kfont(rr, fs, True, WHITE)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            c = tbl.cell(i, j); c.text = ""
            c.fill.solid(); c.fill.fore_color.rgb = LIGHT if i % 2 else WHITE
            rr = c.text_frame.paragraphs[0].add_run(); rr.text = val
            _kfont(rr, fs, j == 0, GRAY)
    return y + cell_h * n


def build(out_path):
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE

    # ---------- 헤더 ----------
    _rect(s, 0, 0, W, MARGIN + HEADER_H, NAVY)
    _rect(s, 0, MARGIN + HEADER_H, W, 0.18, ACCENT)
    _text(s, MARGIN, 0.8, W - 2 * MARGIN, 2.9, [
        [("단일 스마트폰 영상 기반 마커리스 컴퓨터비전을 이용한", 44, True, WHITE)],
        [("야구 투구 동작 분석 및 측정 신뢰도 검증", 44, True, WHITE)],
    ], align=PP_ALIGN.CENTER)
    _text(s, MARGIN, 3.75, W - 2 * MARGIN, 1.4, [
        [("저비용 2D 자세추정의 측정 신뢰도 검증 방법론", 24, False, RGBColor(0xCF, 0xE0, 0xF1))],
        [("김현준 · Sports Technology Club     |     KSEF 2026 / ISEF 2027",
          22, False, RGBColor(0xCF, 0xE0, 0xF1))],
    ], align=PP_ALIGN.CENTER)

    # ================= 열 1 =================
    c0 = COL_X[0]
    y = _panel(s, 0, COL_TOP, 8.6, "1. 연구 배경 — 접근성의 격차")
    _text(s, c0, y, COL_W, 7.4, _bullets([
        "투구는 인체 최고속 관절 운동(어깨 ~7,000°/s), 팔꿈치 외반토크 ~64 N·m → 부상 위험 큼",
        "정량 분석 표준 = 광학식 마커 모션캡처 / 다중카메라(KinaTrax·Hawk-Eye·Theia3D)",
        (1, "고가 장비 · 통제된 실험실 · 전문 인력 필요"),
        "정작 분석이 필요한 학생·유소년·아마추어는 접근 불가",
        (1, "→ 단일 스마트폰으로 이 격차를 좁힐 수 있는가?"),
    ], size=21))

    y2 = _panel(s, 0, COL_TOP + 9.0, 7.8, "2. 연구 목적 · 연구 질문")
    _text(s, c0, y2, COL_W, 6.6, _bullets([
        "요구사항: 두 촬영 각도 영상 각각에 스켈레톤('졸라맨') 생성",
        "RQ1. 서로 다른 조건(주간 측면/역광 정면)에서 안정적 자세추정이 가능한가?",
        "RQ2. 투구 단계·핵심 관절 운동학을 자동 산출할 수 있는가?",
        "RQ3. 참값(모션캡처) 없이 저비용 2D 측정의 신뢰도를 어떻게 검증하는가?",
    ], size=21))

    y3 = _panel(s, 0, COL_TOP + 17.6, 8.4, "3. 관련 연구 비교")
    ny = _mini_table(s, c0 + 0.2, y3, COL_W - 0.4,
        ["시스템", "카메라", "2D/3D", "대상"],
        [
            ["본 연구", "1(폰)", "2D", "학생·아마추어"],
            ["pitchAI", "1", "3D", "연구실·프로"],
            ["Fleisig'24", "9", "3D", "프로"],
            ["Aguinaldo'25", "다수", "3D", "스타디움"],
            ["OpenCap", "2+", "3D", "연구자"],
        ], cell_h=0.72, fs=16)
    _text(s, c0 + 0.2, ny + 0.15, COL_W - 0.4, 1.4, [
        [("2025 J Sports Sci: \"단일 카메라 마커리스의 투구 분석", 16, False, ACCENT)],
        [("타당성 연구는 현재 존재하지 않는다\"", 16, False, ACCENT)],
    ])

    # ================= 열 2 =================
    c1 = COL_X[1]
    y = _panel(s, 1, COL_TOP, 6.3, "4. 방법 — 시스템 구조")
    _text(s, c1, y, COL_W, 5.2, [
        [("입력 영상 → 전처리(역광 CLAHE) → 자세추정 엔진 → One-Euro 스무딩", 19, False, GRAY)],
        [("→ 스켈레톤 오버레이 → 투구 단계 자동 분할 → 관절 운동학 → 검증 통계", 19, False, GRAY)],
        [("· 두 엔진 통일(COCO 17관절): MediaPipe BlazePose(Heavy) + YOLO11-pose", 19, False, GRAY)],
        [("· One-Euro 필터: 릴리스(빠름) 지연 최소, 정지 국면 지터 억제", 19, False, GRAY)],
        [("· 전 과정 Python 3.11 + 오픈소스 → 학생 재현 가능", 19, True, NAVY)],
    ])

    y2 = _panel(s, 1, COL_TOP + 6.7, 5.6, "5. 방법 — 측정 신뢰도 검증")
    _text(s, c1, y2, COL_W, 4.6, _bullets([
        "CV 각도 vs 참값(2인 수동 주석) 대조",
        "Bland-Altman(편향+95% 일치한계), ICC(2,1)+95% CI, RMSE·MAE·r",
        "검사-재검사 / 평가자 간·내 신뢰도",
        (1, "참값 부재 시 → 엔진 교차검증을 신뢰도 대리지표로"),
    ], size=20))

    # 그림: 정면 역광 오버레이
    img_y = COL_TOP + 12.7
    _rect(s, c1, img_y, COL_W, 0.95, BLUE)
    _text(s, c1 + 0.1, img_y + 0.05, COL_W - 0.2, 0.85,
          [[("스켈레톤 오버레이 (역광 정면)", 26, True, WHITE)]], anchor=MSO_ANCHOR.MIDDLE)
    front = os.path.join(FIG, "fig_front_overlay.png")
    if os.path.exists(front):
        s.shapes.add_picture(front, Inches(c1 + COL_W * 0.18), Inches(img_y + 1.2),
                             height=Inches(9.4))
    _text(s, c1, img_y + 10.8, COL_W, 0.8,
          [[("난이도 높은 역광·소형 피사체에서도 해부학적으로 타당한 추적 (197/197=100%)",
             16, False, GRAY)]], align=PP_ALIGN.CENTER)

    # ================= 열 3 =================
    c2 = COL_X[2]
    y = _panel(s, 2, COL_TOP, 3.0, "6. 결과 ① 두 각도 100% 검출")
    _mini_table(s, c2 + 0.2, y, COL_W - 0.4,
        ["영상", "조건", "검출률"],
        [["측면", "주간·순광", "122/122 = 100%"],
         ["정면", "역광·소형", "197/197 = 100%"]], cell_h=0.62, fs=16)

    y2 = COL_TOP + 3.4
    _rect(s, c2, y2, COL_W, 0.95, BLUE)
    _text(s, c2 + 0.1, y2 + 0.05, COL_W - 0.2, 0.85,
          [[("결과 ② 운동학·단계 자동 검출", 26, True, WHITE)]], anchor=MSO_ANCHOR.MIDDLE)
    kin = os.path.join(FIG, "fig_kinematics.png")
    kh = 7.3                      # 높이 기준(원본 10x11 세로형) → 아래 패널과 겹침 방지
    kw = kh * 10 / 11
    if os.path.exists(kin):
        s.shapes.add_picture(kin, Inches(c2 + (COL_W - kw) / 2), Inches(y2 + 1.15),
                             height=Inches(kh))
    _text(s, c2, y2 + 1.15 + kh + 0.05, COL_W, 0.8,
          [[("손목 최고속이 릴리스에 정렬 · 리드무릎 118°→신전 · 몸통 38°(정상범위)",
             15, False, GRAY)]], align=PP_ALIGN.CENTER)

    y3 = _panel(s, 2, COL_TOP + 13.2, 4.4, "결과 ③ 교차검증 · 통계")
    _text(s, c2, y3, COL_W, 3.6, _bullets([
        "두 엔진 일치: 몸통 길이의 6~7% 이내 → 교차검증 가능",
        "MediaPipe: 지터 우위 + 처리속도 약 7배 → 기본 엔진",
        (1, "ICC(2,1) 정합성: Shrout&Fleiss 예제 0.290과 정확히 일치 ✓"),
    ], size=19))

    y4 = _panel(s, 2, COL_TOP + 18.0, 5.2, "7. 결론 · 향후 · 한계")
    _text(s, c2, y4, COL_W, 4.4, [
        [("가치 = 정확도가 아닌 접근성 + 방법론적 엄밀성", 20, True, NAVY)],
        [("• 향후: IRB 하 240fps·2카메라 다투수 데이터셋, 참값 대조 검증", 18, False, GRAY)],
        [("• 향후: 발판 지면반력(GRF)–영상 릴리스 동기화, 다카메라 3D 확장", 18, False, GRAY)],
        [("• 한계: 2D 어깨회전 미산출 · 30fps 과소표집 · 표본 소수(예비)", 18, False, ACCENT)],
    ])

    # 참고문헌(하단 띠)
    _rect(s, MARGIN, H - 2.0, W - 2 * MARGIN, 1.5, LIGHT, line=RGBColor(0xD5, 0xDE, 0xE8))
    _text(s, MARGIN + 0.3, H - 1.9, W - 2 * MARGIN - 0.6, 1.3, [
        [("핵심 참고문헌: ", 15, True, NAVY),
         ("[1] Fleisig et al. 1995, AJSM · [4] Fleisig et al. 2024, Sports Biomech · "
          "[7] Dobos et al. 2025 (pitchAI) · [9] Uhlrich et al. 2023 (OpenCap) · "
          "[12] Bazarevsky et al. 2020 (BlazePose) · [15] Shrout & Fleiss 1979 (ICC)",
          15, False, GRAY)],
    ])

    prs.save(out_path)
    print(f"[poster] 저장: {out_path}  (슬라이드 {len(prs.slides._sldIdLst)}장, {W}x{H}인치)")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("-o", "--output", default="paper/투구분석_포스터.pptx")
    args = ap.parse_args()
    build(args.output)


if __name__ == "__main__":
    main()
