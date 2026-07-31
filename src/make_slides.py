"""발표자료(pptx) 생성 — KSEF/ISEF 발표용 ~15장.

한글 폰트(맑은 고딕), 결과 이미지 임베드. 16:9.

사용: python src/make_slides.py -o paper/투구분석_발표.pptx
"""

from __future__ import annotations

import argparse
import os

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

KOR = "맑은 고딕"
NAVY = RGBColor(0x1F, 0x4E, 0x79)
LIGHT = RGBColor(0xEB, 0xF1, 0xF8)
GRAY = RGBColor(0x44, 0x44, 0x44)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FIG = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "paper", "figs")


def _kfont(run, size, bold=False, color=GRAY):
    run.font.name = KOR
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.makeelement(qn("a:ea"), {"typeface": KOR})
    rPr.append(ea)


def _bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_title_slide(prs, title, subtitle, meta):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, NAVY)
    box = s.shapes.add_textbox(Inches(0.7), Inches(2.1), Inches(11.9), Inches(3))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    _kfont(p.add_run(), 1)  # placeholder
    p.runs[0].text = title
    _kfont(p.runs[0], 30, bold=True, color=WHITE)
    p2 = tf.add_paragraph(); p2.space_before = Pt(14)
    r = p2.add_run(); r.text = subtitle; _kfont(r, 15, color=RGBColor(0xBD, 0xD3, 0xEC))
    p3 = tf.add_paragraph(); p3.space_before = Pt(24)
    r = p3.add_run(); r.text = meta; _kfont(r, 13, color=RGBColor(0xBD, 0xD3, 0xEC))
    return s


def _header(slide, title, idx=None):
    bar = slide.shapes.add_shape(1, 0, 0, Emu(int(13.333 * 914400)), Inches(0.95))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    tf = bar.text_frame; tf.margin_left = Inches(0.5); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; r = p.add_run(); r.text = title
    _kfont(r, 22, bold=True, color=WHITE)
    if idx is not None:
        nb = slide.shapes.add_textbox(Inches(12.2), Inches(0.25), Inches(0.9), Inches(0.5))
        rp = nb.text_frame.paragraphs[0]; rp.alignment = PP_ALIGN.RIGHT
        rr = rp.add_run(); rr.text = str(idx); _kfont(rr, 12, color=RGBColor(0xBD, 0xD3, 0xEC))


def add_bullets(prs, title, bullets, idx=None):
    s = prs.slides.add_slide(prs.slide_layouts[6]); _bg(s, WHITE)
    _header(s, title, idx)
    box = s.shapes.add_textbox(Inches(0.7), Inches(1.3), Inches(11.9), Inches(5.9))
    tf = box.text_frame; tf.word_wrap = True
    for k, b in enumerate(bullets):
        level = 0
        text = b
        if isinstance(b, tuple):
            level, text = b
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.level = level; p.space_after = Pt(9)
        bullet = "•  " if level == 0 else "–  "
        r = p.add_run(); r.text = bullet + text
        _kfont(r, 18 if level == 0 else 15, bold=(level == 0), color=(NAVY if level == 0 else GRAY))
    return s


def add_image_slide(prs, title, img, caption, idx=None, bullets=None):
    s = prs.slides.add_slide(prs.slide_layouts[6]); _bg(s, WHITE)
    _header(s, title, idx)
    path = os.path.join(FIG, img)
    if bullets:
        # 좌: 이미지, 우: 설명.
        if os.path.exists(path):
            s.shapes.add_picture(path, Inches(0.5), Inches(1.3), height=Inches(5.4))
        box = s.shapes.add_textbox(Inches(7.7), Inches(1.5), Inches(5.2), Inches(5.4))
        tf = box.text_frame; tf.word_wrap = True
        for k, b in enumerate(bullets):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.space_after = Pt(10)
            r = p.add_run(); r.text = "•  " + b; _kfont(r, 15, color=GRAY)
    else:
        if os.path.exists(path):
            s.shapes.add_picture(path, Inches(2.6), Inches(1.2), height=Inches(5.2))
    if caption:
        cap = s.shapes.add_textbox(Inches(0.5), Inches(6.85), Inches(12.3), Inches(0.5))
        rp = cap.text_frame.paragraphs[0]; rp.alignment = PP_ALIGN.CENTER
        r = rp.add_run(); r.text = caption; _kfont(r, 11, color=GRAY); r.font.italic = True
    return s


def add_table_slide(prs, title, header, rows, idx=None, note=None):
    s = prs.slides.add_slide(prs.slide_layouts[6]); _bg(s, WHITE)
    _header(s, title, idx)
    nrows, ncols = len(rows) + 1, len(header)
    gt = s.shapes.add_table(nrows, ncols, Inches(0.5), Inches(1.35),
                            Inches(12.3), Inches(0.4 * nrows)).table
    for j, h in enumerate(header):
        c = gt.cell(0, j); c.text = ""
        r = c.text_frame.paragraphs[0].add_run(); r.text = h
        _kfont(r, 12, bold=True, color=WHITE)
        c.fill.solid(); c.fill.fore_color.rgb = NAVY
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            c = gt.cell(i, j); c.text = ""
            r = c.text_frame.paragraphs[0].add_run(); r.text = val
            _kfont(r, 11, bold=(j == 0), color=GRAY)
            c.fill.solid(); c.fill.fore_color.rgb = LIGHT if i % 2 else WHITE
    if note:
        nb = s.shapes.add_textbox(Inches(0.5), Inches(6.9), Inches(12.3), Inches(0.5))
        r = nb.text_frame.paragraphs[0].add_run(); r.text = note
        _kfont(r, 11, color=GRAY); r.font.italic = True
    return s


def build(out_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1. 표지
    add_title_slide(
        prs,
        "단일 스마트폰 영상 기반\n마커리스 컴퓨터비전 야구 투구 분석",
        "저비용 2D 자세추정의 측정 신뢰도 검증 방법론",
        "Sports Technology Club · 김현준  |  KSEF 2026 / ISEF 2027",
    )

    # 2. 배경/문제
    add_bullets(prs, "연구 배경 — 접근성의 격차", [
        "투구는 인체 최고속 관절 운동(어깨 ~7,000°/s), 팔꿈치 외반토크 ~64 N·m → 부상 위험 큼",
        "정량 분석의 표준 = 광학식 마커 모션캡처 / 다중카메라(KinaTrax·Hawk-Eye·Theia3D)",
        (1, "고가 장비 · 통제된 실험실 · 전문 인력 필요"),
        "결과: 정작 분석이 필요한 학생·유소년·아마추어는 접근 불가",
        "→ '저비용 · 단일 스마트폰'으로 이 격차를 좁힐 수 있는가?",
    ], idx=2)

    # 3. 목적
    add_bullets(prs, "연구 목적 및 연구 질문", [
        "최초 요구사항: 두 촬영 각도 영상 각각에 관절 스켈레톤('졸라맨') 생성",
        "RQ1. 단일 스마트폰으로 서로 다른 조건(주간 측면/역광 정면)에서 안정적 자세추정이 가능한가?",
        "RQ2. 투구 단계와 핵심 관절 운동학을 자동 산출할 수 있는가?",
        "RQ3. 참값(모션캡처)이 없는 환경에서 저비용 2D 측정의 신뢰도를 어떻게 검증하는가?",
    ], idx=3)

    # 4. 관련연구 비교
    add_table_slide(prs, "관련 연구 비교",
        ["시스템", "카메라", "비용", "2D/3D", "대상", "오픈"],
        [
            ["본 연구", "1(스마트폰)", "소비자급", "2D", "학생·아마추어", "예"],
            ["pitchAI", "1", "상용", "3D", "연구실·프로", "아니오"],
            ["Fleisig 2024", "9", "실험실", "3D", "프로", "아니오"],
            ["Aguinaldo 2025", "다수", "매우높음", "3D", "스타디움", "아니오"],
            ["PitcherNet", "1(방송)", "고연산", "3D", "프로분석", "아니오"],
            ["OpenCap/Pose2Sim", "2+", "중저", "3D", "연구자", "예"],
        ], idx=4,
        note="2025 J Sports Sci: \"단일 카메라 마커리스의 투구 분석 타당성 연구는 현재 존재하지 않는다\"")

    # 5. 신규성
    add_bullets(prs, "신규성 — 정확도가 아닌 접근성", [
        "본 연구는 실험실 시스템보다 정확하다고 주장하지 않음 (마커 모캡이 여전히 표준)",
        "기여 1 — 접근성: 단일 소비자 스마트폰만으로 동작하는 파이프라인",
        "기여 2 — 엔진 교차검증: 참값 없을 때 독립 두 엔진 일치도를 신뢰도 대리지표로",
        "기여 3 — 재현 가능한 검증 방법론: Bland-Altman · ICC · RMSE",
        "이 조합을 야구 투구에 적용한 선행연구는 확인되지 않음",
    ], idx=5)

    # 6. 시스템 구조
    add_bullets(prs, "시스템 구조", [
        "입력 영상 → 전처리(역광 CLAHE 보정, 선택)",
        (1, "→ 자세추정 엔진(MediaPipe / YOLO11) → COCO 17관절"),
        (1, "→ One-Euro 스무딩(지터 제거) → 스켈레톤 오버레이"),
        (1, "→ 투구 단계 자동 분할 → 관절 운동학 → 검증 통계"),
        "전 과정 Python 3.11 + 오픈소스, 학생 재현 가능",
    ], idx=6)

    # 7. 방법 - 자세추정/스무딩
    add_bullets(prs, "방법 — 자세추정과 스무딩", [
        "두 엔진을 동일 출력 규약(COCO 17관절)으로 통일 → 교체·비교 가능",
        (1, "MediaPipe BlazePose(Heavy): 33랜드마크, CPU 실시간"),
        (1, "YOLO11-pose: 17관절, 역광·소형 대비 백업"),
        "One-Euro 필터: 속도 적응형 → 릴리스(빠름)엔 지연 최소, 정지엔 떨림 억제",
        "다인원 프레임에서 관절신뢰도 합 최대 1인(투수) 자동 선택",
    ], idx=7)

    # 8. 방법 - 운동학/단계
    add_bullets(prs, "방법 — 운동학 및 단계 분할", [
        "관절각 = 세 점 내각; 팔꿈치·무릎은 굴곡각(180°−내각)으로 표기",
        "몸통 기울기 = 골반→어깨 벡터의 연직 대비 각",
        "투구 단계 자동 검출:",
        (1, "다리 들기 최고점 · 앞발 착지 · 릴리스(손목 최고속)"),
        "던지는 팔·리드 다리 = 손목 최고속 기준 자동 판정",
    ], idx=8)

    # 9. 방법 - 검증
    add_bullets(prs, "방법 — 측정 신뢰도 검증", [
        "CV 각도 vs 참값(훈련된 2인 수동 주석) 대조",
        "Bland-Altman: 편향 + 95% 일치한계",
        "ICC(2,1): 이원 임의효과·절대일치 + 95% 신뢰구간 (Shrout & Fleiss 1979)",
        "RMSE · MAE · Pearson r; 검사-재검사/평가자 간·내 신뢰도",
        "참값 부재 시 → 엔진 교차검증을 대리지표로 활용",
    ], idx=9)

    # 10. 결과1 - 검출
    add_image_slide(prs, "결과 1 — 두 각도 100% 검출 (RQ1)",
        "fig_side_overlay.png", "그림. 측면(주간) 영상 스켈레톤 오버레이", idx=10,
        bullets=[
            "측면(주간): 122/122 = 100%",
            "정면(역광·소형): 197/197 = 100%",
            "난이도 높은 역광 정면에서도 해부학적으로 타당한 추적",
            "요구사항(두 각도 졸라맨) 충족",
        ])

    # 11. 결과2 - 운동학
    add_image_slide(prs, "결과 2 — 운동학·단계 자동 검출 (RQ2)",
        "fig_kinematics.png", "그림. 투구 운동학 시계열과 자동 단계 마커", idx=11,
        bullets=[
            "단계: 다리들기(1.17s)→착지(2.10s)→릴리스(2.33s)",
            "손목 최고속(~2,800px/s)이 릴리스에 정확히 정렬",
            "리드 무릎 118° 굴곡→신전(리드레그 블록)",
            "몸통 릴리스 시 전방 38° (정상 28–42° 내)",
        ])

    # 12. 결과3 - 교차검증/ICC
    add_bullets(prs, "결과 3 — 엔진 교차검증 & 통계 검증 (RQ3)", [
        "두 엔진 동일지점 지목: 몸통 길이의 6~7% 이내 일치",
        (1, "측면 13.5px/7%, 정면 7.2px/6% → 상호 교차검증 가능"),
        "MediaPipe: 지터 근소 우위 + 처리속도 약 7배 → 기본 엔진 확정",
        "ICC(2,1) 구현 정합성: Shrout&Fleiss 예제값 0.290과 정확히 일치 ✓",
        "합성 참값 종단 검증 완료 → 실측 주석 확보 시 즉시 실제 검증 가능",
    ], idx=12)

    # 13. 고찰
    add_bullets(prs, "고찰 — 한계와 의의", [
        "한계:",
        (1, "2D 단일카메라 → 횡단면 어깨 회전 측정 불가(의도적 미산출)"),
        (1, "30fps는 릴리스 과소표집 → 정식 촬영 240fps 필요"),
        (1, "참값·표본 부족 → 절대 정확도는 추후 검증 후 결론"),
        "의의: 접근 가능한 파이프라인 실증 + 과학적으로 정직한 재현 가능 검증 방법론",
    ], idx=13)

    # 14. 결론/향후
    add_bullets(prs, "결론 및 향후 연구", [
        "단일 스마트폰으로 100% 검출·자동 운동학 달성; 가치는 접근성 + 방법론적 엄밀성",
        "향후 1 — IRB 승인 하 240fps·2카메라 다투수 데이터셋 구축",
        "향후 2 — 수동 주석 참값 대조로 변수별 타당 범위 확정",
        "향후 3 — 발판 지면반력(GRF)과 영상 릴리스 동기화 → '하체 힘→공' 정량화",
        "향후 4 — Pose2Sim/OpenCap 방식 다중카메라 3D 확장",
    ], idx=14)

    # 15. 마무리
    s = add_bullets(prs, "요약", [
        "요구사항(두 각도 졸라맨) → 학술 연구로 확장 완성",
        "접근성: 단일 스마트폰, 오픈소스, 학생 재현 가능",
        "엄밀성: 엔진 교차검증 + Bland-Altman/ICC/RMSE (ICC 정합성 검증 완료)",
        "정직성: 2D·30fps·어깨회전 한계 명시",
        "감사합니다.",
    ], idx=15)

    prs.save(out_path)
    print(f"[pptx] 저장: {out_path}  (슬라이드 {len(prs.slides.__iter__.__self__._sldIdLst)}장)")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("-o", "--output", default="paper/투구분석_발표.pptx")
    args = ap.parse_args()
    build(args.output)


if __name__ == "__main__":
    main()
