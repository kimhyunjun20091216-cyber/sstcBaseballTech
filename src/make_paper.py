"""논문.md → docx 변환.

프로젝트에서 사용한 마크다운 하위집합을 처리한다:
제목(#/##/###), 문단, **굵게**, 불릿(-), 표(| |), 이미지 ![alt](path),
수평선(---), 참고문헌([n] ...). 한글 폰트(맑은 고딕)를 문서 전역에 적용한다.

사용: python src/make_paper.py paper/논문.md -o paper/투구분석_논문.docx
"""

from __future__ import annotations

import argparse
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

KOR_FONT = "맑은 고딕"      # Windows 기본 한글 폰트(뷰어 호환성 높음)
ACCENT = RGBColor(0x1F, 0x4E, 0x79)


def _set_font(run, name=KOR_FONT):
    run.font.name = name
    r = run._element.rPr.rFonts
    r.set(qn("w:eastAsia"), name)   # 한글(동아시아) 글꼴 지정


def _base_style(doc):
    st = doc.styles["Normal"]
    st.font.name = KOR_FONT
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), KOR_FONT)


def _add_runs(par, text):
    """**굵게** 마크업을 런으로 분해해 추가."""
    for i, seg in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not seg:
            continue
        run = par.add_run(seg)
        _set_font(run)
        if i % 2 == 1:               # 홀수 = ** ** 안쪽
            run.bold = True


def _add_table(doc, rows):
    header = [c.strip() for c in rows[0].strip("|").split("|")]
    body = [[c.strip() for c in r.strip("|").split("|")] for r in rows[2:]]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(header):
        cell = t.rows[0].cells[j]
        cell.paragraphs[0].text = ""
        _add_runs(cell.paragraphs[0], h)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for brow in body:
        cells = t.add_row().cells
        for j, val in enumerate(brow[:len(header)]):
            cells[j].paragraphs[0].text = ""
            _add_runs(cells[j].paragraphs[0], val)
            for run in cells[j].paragraphs[0].runs:
                run.font.size = Pt(9)
    doc.add_paragraph()


def convert(md_path, out_path):
    with open(md_path, encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()
    _base_style(doc)
    base_dir = os.path.dirname(os.path.abspath(md_path))

    i = 0
    n = len(lines)
    while i < n:
        raw = lines[i].rstrip("\n")
        line = raw.strip()

        # 표: | ... | 로 시작하고 다음 줄이 |---| 구분선.
        if line.startswith("|") and i + 1 < n and set(lines[i + 1].strip()) <= set("|-: "):
            block = []
            while i < n and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            _add_table(doc, block)
            continue

        # 이미지: ![alt](path)
        m = re.match(r"!\[(.*?)\]\((.*?)\)", line)
        if m:
            alt, path = m.group(1), m.group(2)
            img = os.path.join(base_dir, path)
            if os.path.exists(img):
                doc.add_picture(img, width=Inches(5.6))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = cap.add_run(alt)
                _set_font(run)
                run.font.size = Pt(9)
                run.italic = True
            i += 1
            continue

        if not line:
            i += 1
            continue

        if line == "---":
            i += 1
            continue

        # 제목.
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line[2:]); _set_font(run)
            run.bold = True; run.font.size = Pt(16); run.font.color.rgb = ACCENT
        elif line.startswith("## "):
            doc.add_paragraph()
            p = doc.add_paragraph()
            run = p.add_run(line[3:]); _set_font(run)
            run.bold = True; run.font.size = Pt(13.5); run.font.color.rgb = ACCENT
        elif line.startswith("### "):
            p = doc.add_paragraph()
            run = p.add_run(line[4:]); _set_font(run)
            run.bold = True; run.font.size = Pt(11.5)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            _add_runs(p, line[2:])
        else:
            p = doc.add_paragraph()
            _add_runs(p, line)
            # 영문 부제/저자 줄은 가운데 정렬(첫 문단 근처).
            if line.startswith("**A ") or "Sports Technology Club" in line or "KSEF" in line:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        i += 1

    doc.save(out_path)
    print(f"[docx] 저장: {out_path}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("md")
    ap.add_argument("-o", "--output", default="paper/투구분석_논문.docx")
    args = ap.parse_args()
    convert(args.md, args.output)


if __name__ == "__main__":
    main()
